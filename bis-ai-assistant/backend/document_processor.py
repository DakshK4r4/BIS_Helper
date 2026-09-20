import os
import io
import shutil
from pypdf import PdfReader
from docx import Document
from PIL import Image
import pytesseract

# Configure Tesseract path if available in common locations
TESSERACT_CANDIDATES = [
    shutil.which("tesseract"),
    "/opt/homebrew/bin/tesseract",
    "/usr/local/bin/tesseract",
    "/usr/bin/tesseract",
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
]

TESSERACT_AVAILABLE = False
for candidate in TESSERACT_CANDIDATES:
    if candidate and os.path.exists(candidate):
        pytesseract.pytesseract.tesseract_cmd = candidate
        TESSERACT_AVAILABLE = True
        break


def extract_pdf_text(path):
    """
    Extract selectable text from PDF using pypdf.
    """
    reader = PdfReader(path)
    text_parts = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text)
    return "\n".join(text_parts).strip()


def extract_pdf_ocr(path):
    """
    Fallback OCR for scanned PDFs: extract images embedded in pages and run OCR.
    """
    if not TESSERACT_AVAILABLE:
        return ""

    try:
        reader = PdfReader(path)
        ocr_texts = []
        for page_idx, page in enumerate(reader.pages):
            for img_file in page.images:
                try:
                    img_bytes = img_file.data
                    img = Image.open(io.BytesIO(img_bytes))
                    txt = pytesseract.image_to_string(img)
                    if txt.strip():
                        ocr_texts.append(txt.strip())
                except Exception as e:
                    print(f"Notice: OCR on page {page_idx} image failed: {e}")
        return "\n".join(ocr_texts).strip()
    except Exception as e:
        print(f"Scanned PDF OCR error: {e}")
        return ""


def extract_docx_text(path):
    """
    Extract paragraphs and tables from DOCX.
    """
    document = Document(path)
    text_parts = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text.strip())

    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_parts.append(row_text)

    return "\n".join(text_parts).strip()


def ocr_image(path):
    """
    Perform OCR on PNG, JPG, JPEG images.
    """
    if not TESSERACT_AVAILABLE:
        # Fallback when system tesseract is not installed: read basic image metadata
        try:
            with Image.open(path) as img:
                return f"[Image Document: {os.path.basename(path)} | Format: {img.format} | Dimensions: {img.width}x{img.height} | Mode: {img.mode}. System Tesseract binary is required for raw pixel text extraction.]"
        except Exception:
            return f"[Image Document: {os.path.basename(path)}]"

    try:
        image = Image.open(path)
        # Convert to RGB if palette/RGBA
        if image.mode not in ("L", "RGB"):
            image = image.convert("RGB")
        text = pytesseract.image_to_string(image)
        return text.strip()
    except Exception as e:
        print(f"Image OCR error for {path}: {e}")
        return f"[Image Document: {os.path.basename(path)} — OCR extraction encountered error: {str(e)}]"


def process_document(path):
    """
    Unified active document processor for PDF, DOCX, and images.
    Returns:
    {
        "text": extracted text,
        "ocr_used": bool,
        "file_type": str,
        "status": "PROCESSED" | "OCR_REQUIRED" | "OCR_APPLIED"
    }
    """
    if not os.path.exists(path):
        raise FileNotFoundError(f"Document file not found at: {path}")

    extension = os.path.splitext(path)[1].lower()

    # 1. PDF
    if extension == ".pdf":
        text = extract_pdf_text(path)
        ocr_used = False

        # If selectable text is under 50 characters, test for scanned PDF
        if len(text.strip()) < 50:
            scanned_text = extract_pdf_ocr(path)
            if scanned_text and len(scanned_text.strip()) >= 50:
                text = scanned_text
                ocr_used = True
                status = "OCR_APPLIED"
            else:
                # Scanned PDF without Tesseract OCR available or no text found
                status = "OCR_REQUIRED"
        else:
            status = "PROCESSED"

        return {
            "text": text,
            "ocr_used": ocr_used,
            "file_type": "PDF",
            "status": status
        }

    # 2. DOCX
    elif extension == ".docx":
        text = extract_docx_text(path)
        return {
            "text": text,
            "ocr_used": False,
            "file_type": "DOCX",
            "status": "PROCESSED"
        }

    # 3. Images (PNG, JPG, JPEG)
    elif extension in [".png", ".jpg", ".jpeg"]:
        text = ocr_image(path)
        return {
            "text": text,
            "ocr_used": True if TESSERACT_AVAILABLE else False,
            "file_type": "IMAGE",
            "status": "PROCESSED"
        }

    else:
        raise ValueError(f"Unsupported file format: {extension}. Supported formats: PDF, DOCX, PNG, JPG, JPEG.")